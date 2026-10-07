"""Expense tracker: Capital One alert emails in, spending dashboard out.

Gmail is never touched by this app. A Google Apps Script in Aege's own account
watches the label and POSTs each new email to /ingest.
"""

import csv
import io
import logging
import os
import uuid
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

from flask import Flask, jsonify, request, send_from_directory
from flask_cors import CORS
from sqlalchemy import func, select

import categorize
import parser as email_parser
from db import IngestLog, MerchantRule, SessionLocal, Transaction, init_db

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")
logger = logging.getLogger(__name__)

app = Flask(__name__, static_folder=None)
CORS(app)

STATIC_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static")

# Every "today" and "this month" in the app is this clock, not the server's.
TZ = ZoneInfo(os.environ.get("TZ", "America/Chicago"))

ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "")
INGEST_SECRET = os.environ.get("INGEST_SECRET", "")


def today() -> date:
    return datetime.now(TZ).date()


def password_ok() -> bool:
    if not ADMIN_PASSWORD:
        return True                       # nothing set: the app is open
    sent = request.headers.get("X-Admin-Password") or request.args.get("key") or ""
    return sent == ADMIN_PASSWORD


def require_auth():
    if not password_ok():
        return jsonify({"error": "Unauthorized"}), 401
    return None


# --- pages --------------------------------------------------------------

@app.route("/")
def index():
    return send_from_directory(STATIC_DIR, "index.html")


@app.route("/static/<path:filename>")
def static_files(filename: str):
    return send_from_directory(STATIC_DIR, filename)


@app.route("/health")
def health():
    return jsonify({"status": "ok"})


@app.route("/auth-check")
def auth_check():
    return jsonify({"ok": password_ok(), "required": bool(ADMIN_PASSWORD)})


@app.route("/categories")
def list_categories():
    return jsonify(categorize.CATEGORIES)


# --- ingest -------------------------------------------------------------

def _log(session, gmail_id: str, subject: str, status: str, reason: str) -> None:
    existing = session.scalar(select(IngestLog).where(IngestLog.gmail_id == gmail_id))
    if existing:
        existing.status = status
        existing.reason = reason
        existing.subject = subject
    else:
        session.add(IngestLog(gmail_id=gmail_id, subject=subject, status=status, reason=reason))


@app.route("/ingest", methods=["POST"])
def ingest():
    """One email from the Apps Script. Parse it, file it, and say what happened.

    Always answers 200 for anything we understood well enough to record — the
    script has already moved on, and a failure is worth keeping in the log
    rather than retrying forever.
    """
    if INGEST_SECRET and request.headers.get("X-Ingest-Secret") != INGEST_SECRET:
        return jsonify({"error": "Unauthorized"}), 401

    data = request.get_json(silent=True) or {}
    gmail_id = (data.get("gmail_id") or "").strip()
    if not gmail_id:
        return jsonify({"error": "gmail_id is required"}), 400

    subject = (data.get("subject") or "").strip()
    body = data.get("body") or ""

    received = today()
    raw_date = (data.get("date") or "").strip()
    if raw_date:
        try:
            received = datetime.fromisoformat(raw_date.replace("Z", "+00:00")).astimezone(TZ).date()
        except ValueError:
            pass

    with SessionLocal() as session:
        # The script can hand us the same message twice; the first answer stands.
        if session.scalar(select(Transaction).where(Transaction.gmail_id == gmail_id)):
            return jsonify({"status": "duplicate", "reason": "already logged"})

        result = email_parser.parse_email(subject, body, received)

        if result.status != "transaction":
            _log(session, gmail_id, subject, result.status, result.reason)
            session.commit()
            return jsonify({"status": result.status, "reason": result.reason})

        txn = result.txn
        category = categorize.categorize(session, txn["merchant"])
        row = Transaction(
            gmail_id=gmail_id,
            date=txn["date"],
            merchant=txn["merchant"],
            merchant_raw=txn["merchant_raw"],
            amount=txn["amount"],
            category=category,
            card_last4=txn.get("card_last4"),
            source="email",
        )
        session.add(row)
        _log(session, gmail_id, subject, "inserted", f"{txn['merchant']} ${txn['amount']:.2f}")
        session.commit()
        logger.info("Logged %s %s for %s", txn["merchant"], txn["amount"], txn["date"])
        return jsonify({"status": "inserted", "transaction": row.to_dict()})


@app.route("/ingest-log")
def ingest_log():
    if (err := require_auth()):
        return err
    limit = min(int(request.args.get("limit", 50)), 200)
    with SessionLocal() as session:
        rows = session.scalars(
            select(IngestLog).order_by(IngestLog.id.desc()).limit(limit)
        ).all()
        return jsonify([r.to_dict() for r in rows])


# --- the numbers --------------------------------------------------------

def month_bounds(month: str | None) -> tuple[date, date, str]:
    """First and last day of the month being looked at."""
    now = today()
    if month:
        try:
            year, mon = (int(p) for p in month.split("-")[:2])
            first = date(year, mon, 1)
        except (ValueError, TypeError):
            first = now.replace(day=1)
    else:
        first = now.replace(day=1)
    last = (first.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)
    return first, last, first.strftime("%Y-%m")


def compute_summary(session, month: str | None) -> dict:
    """Every number the dashboard shows, worked out in one place.

    The tiles, the chart and the day list all read from this, so they cannot
    disagree with each other.
    """
    first, last, label = month_bounds(month)
    now = today()

    rows = session.scalars(
        select(Transaction)
        .where(Transaction.date >= first, Transaction.date <= last)
        .order_by(Transaction.date.desc(), Transaction.id.desc())
    ).all()

    month_total = sum(float(r.amount) for r in rows)
    today_total = sum(float(r.amount) for r in rows if r.date == now)

    by_category: dict[str, dict] = {}
    by_day: dict[str, dict] = {}
    for r in rows:
        cat = by_category.setdefault(r.category, {"category": r.category, "total": 0.0, "count": 0})
        cat["total"] += float(r.amount)
        cat["count"] += 1
        day = by_day.setdefault(r.date.isoformat(), {"date": r.date.isoformat(), "total": 0.0, "count": 0})
        day["total"] += float(r.amount)
        day["count"] += 1

    # Only count days that have actually happened: a month in progress should
    # not be averaged over days that have not arrived yet.
    in_this_month = first <= now <= last
    days_elapsed = now.day if in_this_month else (last - first).days + 1
    days_in_month = (last - first).days + 1
    daily_avg = month_total / days_elapsed if days_elapsed else 0.0

    return {
        "month": label,
        "month_label": first.strftime("%B %Y"),
        "today": now.isoformat(),
        "today_total": round(today_total, 2),
        "month_total": round(month_total, 2),
        "days_elapsed": days_elapsed,
        "days_in_month": days_in_month,
        "count": len(rows),
        "daily_avg": round(daily_avg, 2),
        "projected_month": round(daily_avg * days_in_month, 2) if in_this_month else round(month_total, 2),
        "is_current_month": in_this_month,
        "by_category": sorted(by_category.values(), key=lambda c: c["total"], reverse=True),
        "by_day": sorted(by_day.values(), key=lambda d: d["date"], reverse=True),
    }


@app.route("/summary")
def summary():
    if (err := require_auth()):
        return err
    with SessionLocal() as session:
        return jsonify(compute_summary(session, request.args.get("month")))


@app.route("/transactions")
def list_transactions():
    if (err := require_auth()):
        return err
    with SessionLocal() as session:
        query = select(Transaction)
        day = request.args.get("date")
        month = request.args.get("month")
        if day:
            try:
                target = datetime.strptime(day, "%Y-%m-%d").date()
            except ValueError:
                return jsonify({"error": "date must be YYYY-MM-DD"}), 400
            query = query.where(Transaction.date == target)
        elif month:
            first, last, _ = month_bounds(month)
            query = query.where(Transaction.date >= first, Transaction.date <= last)
        rows = session.scalars(
            query.order_by(Transaction.date.desc(), Transaction.id.desc())
        ).all()
        return jsonify([r.to_dict() for r in rows])


@app.route("/transactions", methods=["POST"])
def add_transaction():
    """A cash purchase, or anything the card did not email about."""
    if (err := require_auth()):
        return err
    data = request.get_json(silent=True) or {}

    merchant = (data.get("merchant") or "").strip()
    if not merchant:
        return jsonify({"error": "Merchant is required"}), 400
    try:
        amount = round(float(data.get("amount")), 2)
    except (TypeError, ValueError):
        return jsonify({"error": "Amount must be a number"}), 400

    raw_date = (data.get("date") or "").strip()
    try:
        when = datetime.strptime(raw_date, "%Y-%m-%d").date() if raw_date else today()
    except ValueError:
        return jsonify({"error": "Date must be YYYY-MM-DD"}), 400

    with SessionLocal() as session:
        category = (data.get("category") or "").strip()
        if category not in categorize.CATEGORIES:
            category = categorize.categorize(session, merchant)
        row = Transaction(
            gmail_id=f"manual:{uuid.uuid4()}",
            date=when,
            merchant=merchant,
            merchant_raw=merchant,
            amount=amount,
            category=category,
            source="manual",
            note=(data.get("note") or "").strip() or None,
        )
        session.add(row)
        session.commit()
        return jsonify({"success": True, "transaction": row.to_dict()})


@app.route("/transactions/<int:txn_id>", methods=["PATCH"])
def update_transaction(txn_id: int):
    """Correct one. Changing the category also teaches the merchant rule."""
    if (err := require_auth()):
        return err
    data = request.get_json(silent=True) or {}

    with SessionLocal() as session:
        row = session.get(Transaction, txn_id)
        if row is None:
            return jsonify({"error": "Transaction not found"}), 404

        if "merchant" in data:
            merchant = (data.get("merchant") or "").strip()
            if not merchant:
                return jsonify({"error": "Merchant is required"}), 400
            row.merchant = merchant
        if "amount" in data:
            try:
                row.amount = round(float(data.get("amount")), 2)
            except (TypeError, ValueError):
                return jsonify({"error": "Amount must be a number"}), 400
        if "date" in data:
            try:
                row.date = datetime.strptime((data.get("date") or "").strip(), "%Y-%m-%d").date()
            except ValueError:
                return jsonify({"error": "Date must be YYYY-MM-DD"}), 400
        if "note" in data:
            row.note = (data.get("note") or "").strip() or None
        if "category" in data:
            category = (data.get("category") or "").strip()
            if category not in categorize.CATEGORIES:
                return jsonify({"error": "Unknown category"}), 400
            row.category = category
            # The correction is the point: this merchant is filed here from now on.
            categorize.remember(session, row.merchant, category)

        session.commit()
        return jsonify({"success": True, "transaction": row.to_dict()})


@app.route("/transactions/<int:txn_id>", methods=["DELETE"])
def delete_transaction(txn_id: int):
    if (err := require_auth()):
        return err
    with SessionLocal() as session:
        row = session.get(Transaction, txn_id)
        if row is None:
            return jsonify({"error": "Transaction not found"}), 404
        session.delete(row)
        session.commit()
        return jsonify({"success": True})


@app.route("/rules")
def list_rules():
    if (err := require_auth()):
        return err
    with SessionLocal() as session:
        rows = session.scalars(select(MerchantRule).order_by(MerchantRule.pattern)).all()
        return jsonify([r.to_dict() for r in rows])


# --- statement import ---------------------------------------------------

def _read_rows(filename: str, raw: bytes) -> list[dict]:
    """Rows out of a Capital One export, CSV or Excel."""
    if filename.lower().endswith((".xlsx", ".xls")):
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = book.active
        rows = sheet.iter_rows(values_only=True)
        header = [str(h or "").strip() for h in next(rows, [])]
        return [dict(zip(header, [("" if v is None else v) for v in row])) for row in rows]

    text = raw.decode("utf-8-sig", errors="replace")
    return list(csv.DictReader(io.StringIO(text)))


def _pick(row: dict, *names: str) -> str:
    """A column by any of its likely names, ignoring case and spacing."""
    normalised = {str(k or "").strip().lower().replace(" ", ""): v for k, v in row.items()}
    for name in names:
        value = normalised.get(name.lower().replace(" ", ""))
        if value not in (None, ""):
            return str(value).strip()
    return ""


def _statement_date(raw: str) -> date | None:
    raw = (raw or "").strip()
    if not raw:
        return None
    if " " in raw:                                  # Excel gives back a datetime
        raw = raw.split(" ")[0]
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%d/%m/%Y"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            pass
    return None


def _money(raw: str) -> float | None:
    raw = (raw or "").strip().replace("$", "").replace(",", "")
    if not raw:
        return None
    negative = raw.startswith("(") and raw.endswith(")")
    raw = raw.strip("()")
    try:
        value = round(float(raw), 2)
    except ValueError:
        return None
    return -value if negative else value


@app.route("/import", methods=["POST"])
def import_statement():
    """A downloaded Capital One statement, CSV or XLSX.

    Capital One exports put spending in Debit and refunds in Credit, so a row
    has one or the other. Payments to the card are not spending and are left
    out, the same as the payment emails are.
    """
    if (err := require_auth()):
        return err

    upload = request.files.get("file")
    if upload is None or not upload.filename:
        return jsonify({"error": "No file"}), 400

    try:
        rows = _read_rows(upload.filename, upload.read())
    except Exception as exc:                        # noqa: BLE001
        logger.exception("Could not read the statement")
        return jsonify({"error": f"Could not read that file: {exc}"}), 400

    imported = skipped = ignored = 0
    seen: dict[str, int] = {}

    with SessionLocal() as session:
        for row in rows:
            merchant = _pick(row, "Description", "Merchant", "Payee", "Name")
            when = _statement_date(
                _pick(row, "Transaction Date", "TransactionDate", "Date", "Posted Date")
            )
            if not merchant or when is None:
                continue

            debit = _money(_pick(row, "Debit", "Amount Debit", "Withdrawal"))
            credit = _money(_pick(row, "Credit", "Amount Credit", "Deposit"))
            if debit:
                amount = abs(debit)
            elif credit:
                amount = -abs(credit)
            else:
                amount = _money(_pick(row, "Amount"))
                if amount is None:
                    continue

            # A payment to the card is not spending. Refunds from a merchant
            # still come through as credits — only the payment lines go.
            if email_parser.looks_like_card_payment(merchant) or \
                    email_parser.ignored_reason(merchant, merchant):
                ignored += 1
                continue

            # The key doubles as the dedupe: the same date, merchant and amount
            # in the same position is the same transaction. A genuine second
            # identical charge in one day keeps its own slot through the count.
            base = f"stmt:{when.isoformat()}:{merchant.upper()[:60]}:{amount:.2f}"
            seen[base] = seen.get(base, 0) + 1
            key = f"{base}:{seen[base]}"

            if session.scalar(select(Transaction).where(Transaction.gmail_id == key)):
                skipped += 1
                continue
            # An emailed alert for the same purchase arrived first: same day,
            # same amount, so do not log it twice.
            twin = session.scalar(
                select(Transaction).where(
                    Transaction.date == when,
                    Transaction.amount == amount,
                    Transaction.source == "email",
                )
            )
            if twin is not None:
                skipped += 1
                continue

            statement_category = _pick(row, "Category")
            category = categorize.rule_for(session, merchant)
            if not category:
                if statement_category in categorize.CATEGORIES:
                    category = statement_category
                    categorize.remember(session, merchant, category)
                else:
                    category = categorize.categorize(session, merchant)

            session.add(Transaction(
                gmail_id=key,
                date=when,
                merchant=merchant[:200],
                merchant_raw=merchant[:300],
                amount=amount,
                category=category,
                card_last4=_pick(row, "Card No.", "CardNo", "Card")[-4:] or None,
                source="statement",
            ))
            imported += 1

        session.commit()

    logger.info("Statement: imported %s, skipped %s, ignored %s", imported, skipped, ignored)
    return jsonify({"success": True, "imported": imported, "skipped": skipped, "ignored": ignored})


if __name__ == "__main__":
    init_db()
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)))
